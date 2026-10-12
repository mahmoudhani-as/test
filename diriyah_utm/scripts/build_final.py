"""Final verification workbook: before/after %, every live URL with its status now, changes applied,
manual list (with ad names), items waiting on the Whathead reconnect, flagged items, 7-Oct sheet audit, rules."""
import sys, os, json, re, glob, collections
from urllib.parse import urlsplit, parse_qsl
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import FormulaRule
from nomen import Nomen
from derive import PLATFORM_SOURCES

W = sys.argv[1]            # scratchpad dir
OUT = sys.argv[2]
S = json.load(open(f'{W}/work/scored.json'))
SHEET = json.load(open(f'{W}/work/ads.json'))
P1 = json.load(open(f'{W}/work/apply_plan.json'))
P2 = json.load(open(f'{W}/work/apply_plan_legacy.json'))
FLAG = json.load(open(f'{W}/work/legacy_flagged.json'))
MANUAL = json.load(open(f'{W}/work/manual_rows.json'))
BATCHES = json.load(open(f'{W}/apply2/batches.json'))
N = Nomen(glob.glob(f'{W}/in/e2e*.xlsx')[0])

def checks(r, url):
    errs = N.check(url)
    p = r['platform']
    q = dict((k.lower(), v) for k, v in re.findall(r'[?&]([^=&#]+)=([^&#]*)', url))
    src = q.get('utm_source', '').lower()
    if src and p in PLATFORM_SOURCES and src not in PLATFORM_SOURCES[p]:
        errs.append(('SOURCE_PLATFORM_MISMATCH', f'utm_source "{src}" on a {p} ad'))
    if p == 'google' and 'ritz-carlton-signature' in q.get('utm_campaign', '').lower() and not q.get('utm_content', '').lower().startswith('78f964bb_rcsr-search_search_'):
        errs.append(('RITZ_SEARCH_CODE', 'Ritz Google search must use 78f964bb_rcsr-search_search_{lang}'))
    return [e for e in errs if e[0] != 'LP_LANG_MISMATCH']

# ---------- results of live edits ----------
fixed, attempts = {}, {}
for f in sorted(glob.glob(f'{W}/apply/out_*.jsonl')) + sorted(glob.glob(f'{W}/apply2/out_*.jsonl')):
    run = 'Generator fixes (146)' if '/apply/' in f else 'Legacy rebuild'
    for line in open(f):
        if not line.strip(): continue
        a = json.loads(line); a['_run'] = run; a['_file'] = os.path.basename(f)
        for aid in a.get('ad_ids') or []:
            k = (a['platform'], str(aid), a.get('slot'))
            if a['result'] == 'Fixed': fixed[k] = a
            else: attempts[k] = a

# ---------- current-rule targets, manual (user) and planned-but-not-applied ----------
ITEMS = json.load(open(f'{W}/work/legacy_clash_items.json'))
item_of = {(i['platform'], str(i.get('ad_id')), i.get('slot'), i.get('url')): i['expected_url'] for i in ITEMS}
TGT = {}                                      # (platform, ad_id, slot) -> URL under the current rules (legacy plan units)
plan_unit = {}
for u in P2['plan']:
    for ch in u['changes']:
        for aid in u['ad_ids']:
            TGT[(u['platform'], str(aid), ch[0])] = ch[4]; plan_unit[(u['platform'], str(aid), ch[0])] = u
batch_of = {}
for P, fs in BATCHES.items():
    for fn in fs:
        for u in json.load(open(f'{W}/apply2/{fn}')):
            for ch in u['changes']:
                for aid in u['ad_ids']:
                    batch_of[(P, str(aid), ch['slot'])] = fn
USER_BATCHES = {f'in_meta_{i}.json' for i in range(1, 8)} | {f'in_x_{i}.json' for i in range(1, 7)}
user_unit = set()
for fn in USER_BATCHES:
    P = fn.split('_')[1]
    for u in json.load(open(f'{W}/apply2/{fn}')): user_unit.add((P, str(u['unit'])))
manual = {}
for m in MANUAL:                                  # first-plan shared creatives (generator-derived, unchanged)
    if m['why'].startswith('Shared creative/card: ads in different campaigns'):
        manual[(m['platform'], str(m['ad_id']), m['where'].split(' | ', 1)[1])] = m
skipped_unit = {}
for u in P2['skipped']:
    for aid in u['ad_ids']: skipped_unit[(u['platform'], str(aid))] = u
TRIM_ADS = {'1825846574224498', '1825846579472434', '1825848985421841', '1825848991400002', '1825844845218882', '1825846957500434', '1840139996127298'}
META_PARTIAL = {'120253320126610058': '4372459552897444'}
flag_of = {(f['platform'], str(f.get('ad_id')), f.get('slot'), f.get('url')): f['flag'] for f in FLAG}
MANUAL_WHY = 'You asked to do these by hand (the live run for this batch was blocked by the session safety check)'

OFF = re.compile(r'PAUSED|DISABLE|INVALID|NOT_SERVING|REJECTED|DONE|ARCHIVED|REMOVED|ENDED|COMPLETED|DRAFT|HOLD', re.I)
def is_on(s):
    for k in ('campaign_status', 'adset_status', 'ad_status'):
        v = s.get(k)
        if v is not None and not re.match(r'\s*(ACTIVE|ENABLE|ENABLED|SERVING)', str(v), re.I): return False
    return not OFF.search(str(s.get('effective_status') or ''))

# ---------- one row per live slot ----------
rows_src, seen = [], set()
for s in S:
    kind = 'Sitelink' if 'sitelink' in str(s.get('slot') or '').lower() else 'Ad'
    if kind == 'Sitelink':
        k = (s['platform'], s.get('account_id'), s.get('ad_id'), s.get('url'))
        if k in seen: continue
        seen.add(k)
    s['_kind'] = kind; s['_on'] = is_on(s); rows_src.append(s)
rows_src.sort(key=lambda s: (s['platform'], str(s.get('account_name')), s['_kind'], str(s.get('campaign_name')), str(s.get('adset_name')), str(s.get('ad_name')), str(s.get('ad_id')), str(s.get('slot'))))

def status_of(s):
    """-> (status, next step / reason, target URL, live URL after, after verdict)"""
    k = (s['platform'], str(s.get('ad_id')), s.get('slot'))
    if s['verdict'] == 'NO URL': return 'No URL', s.get('notes') or '', '', '', 'NO URL'
    if s['verdict'] == 'CORRECT': return 'Already correct', '', '', s['url'], 'CORRECT'
    tgt = TGT.get(k)
    flag = flag_of.get((s['platform'], str(s.get('ad_id')), s.get('slot'), s.get('url')))
    if k in fixed:
        a = fixed[k]
        if a['_run'].startswith('Generator') or a['after_live'] == tgt:
            e = checks(s, a['after_live'])
            return 'Fixed live', f"{a['_run']}; read back after the write and matched", a['after_live'], a['after_live'], 'WRONG' if e else 'CORRECT'
        if tgt:
            return 'After reconnect', ('Changed live earlier today, but with an older version of the rebuild (it kept "googlesearch"/"tactical"/"rsa1" or the "aud"/"ret" audience tag, or the wrong format). '
                                       'Needs a second edit to the URL in "Correct URL".'), tgt, a['after_live'], 'WRONG'
        return 'Flagged - not fixable', f"Changed live earlier today to the URL in column Q, but under the final rule {flag or 'the old creative token does not match this ad name'} - tell me the creative name to use, or set it yourself.", '', a['after_live'], 'WRONG'
    if k in manual:
        m = manual[k]
        return 'Manual - you', m['why'], m['new'], '', 'WRONG'
    su = skipped_unit.get((s['platform'], str(s.get('ad_id'))))
    if su and not tgt:
        own = item_of.get((s['platform'], str(s.get('ad_id')), s.get('slot'), s.get('url')))
        if own: return 'Manual - you', 'Shared creative/card: ' + su['skip_reason'] + ' - split the creative, then set this ad to this URL', own, '', 'WRONG'
        return 'Flagged - not fixable', (flag or 'not rebuildable') + f" (shared creative/card: {su['skip_reason']})", '', '', 'WRONG'
    if tgt:
        u = plan_unit[k]
        if (s['platform'], str(u['unit'])) in user_unit: return 'Manual - you', MANUAL_WHY, tgt, '', 'WRONG'
        if str(s.get('ad_id')) in TRIM_ADS:
            return 'Manual - you', 'TikTok Spark Ad whose ad text ends with a space: the gateway trims ad text on update, so the API cannot change only the URL. Edit the destination URL in Ads Manager.', tgt, '', 'WRONG'
        if str(s.get('ad_id')) in META_PARTIAL:
            return 'After reconnect', f"New creative {META_PARTIAL[str(s.get('ad_id'))]} with the fixed URL was created but not yet attached to the ad; attach it after reconnect (or edit the URL in Ads Manager).", tgt, '', 'WRONG'
        a = attempts.get(k); d = (a or {}).get('detail') or ''
        if 'OUTCOME UNKNOWN' in d:
            return 'Outcome unknown', 'Write was sent as the sign-in expired; no response. Re-read before any retry.', tgt, '', 'WRONG'
        note = 'Not applied yet: Whathead sign-in expired before this item.'
        if 'Spark' in d or 'BC_AUTH_TT' in d:
            note = 'TikTok Spark Ad: first update method was refused by TikTok (40002); a working method was found on ad 1826384400556033 - will be retried after reconnect. You can also edit it in Ads Manager.'
        return 'After reconnect', note, tgt, '', 'WRONG'
    if flag: return 'Flagged - not fixable', flag + ' (nothing invented)', '', '', 'WRONG'
    if s['derive_status'] == 'FIX':
        return 'Flagged - not fixable', 'Generator URL derived but the creative/card is shared with ads that need a different URL', s.get('expected_url') or '', '', 'WRONG'
    r = (s['derive_reasons'] or ['not in generator'])[0]
    if s['_kind'] == 'Sitelink': r = 'Google sitelink: the generator has no creative code for sitelinks'
    return 'Flagged - not fixable', r, '', '', 'WRONG'

H = ['Platform', 'Account', 'Campaign', 'Ad set / ad group', 'Ad name', 'Ad ID', 'Kind', 'Delivery status', 'Switched on',
     'URL slot', 'Live URL before fix (11-Oct)', 'Verdict before', 'What was wrong', 'Status now', 'Next step / reason',
     'Correct URL (generator)', 'Live URL after fix (read back)', 'Verdict now', 'First row of ad', 'Ad verdict before', 'Ad verdict now']
C = {h: get_column_letter(i) for i, h in enumerate(H, 1)}
recs = []
for s in rows_src:
    st, why, target, after, vnow = status_of(s)
    recs.append((s, st, why, target, after, vnow))
ad_before, ad_now = collections.defaultdict(list), collections.defaultdict(list)
for s, st, why, target, after, vnow in recs:
    k = (s['platform'], s.get('account_id'), str(s.get('ad_id')))
    ad_before[k].append(s['verdict']); ad_now[k].append(vnow)
def agg(v):
    if 'WRONG' in v: return 'WRONG'
    if all(x == 'NO URL' for x in v): return 'NO URL'
    return 'CORRECT'
viol = lambda v: '\n'.join(m for _, m in v)

F = 'Arial'
HDR = Font(name=F, bold=True, color='FFFFFF'); HFILL = PatternFill('solid', fgColor='1F3864')
BODY = Font(name=F, size=10); BOLD = Font(name=F, size=10, bold=True)
TITLE = Font(name=F, size=14, bold=True); H2 = Font(name=F, size=12, bold=True); NOTE = Font(name=F, size=9, italic=True, color='595959')
WRAP = Alignment(wrap_text=True, vertical='top'); TOP = Alignment(vertical='top')
RED = PatternFill('solid', fgColor='F8CBAD'); GREEN = PatternFill('solid', fgColor='C6EFCE'); AMBER = PatternFill('solid', fgColor='FFE699'); BLUE = PatternFill('solid', fgColor='DDEBF7')

def table(ws, headers, rows, name, widths, wrap_cols=()):
    ws.append(headers); hr = ws.max_row
    for c in ws[hr]: c.font = HDR; c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical='center')
    for r in rows:
        ws.append(r)
        for i, c in enumerate(ws[ws.max_row], 1): c.font = BODY; c.alignment = WRAP if i in wrap_cols else TOP
    last = max(ws.max_row, hr + 1)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    t = Table(displayName=name, ref=f'A{hr}:{get_column_letter(len(headers))}{last}')
    t.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True)
    ws.add_table(t); ws.freeze_panes = ws.cell(hr + 1, 1)
    return hr, last

wb = Workbook()
# ---------------- Every URL (live) ----------------
ws = wb.active; ws.title = 'Every URL (live)'
rows, firsts = [], set()
for s, st, why, target, after, vnow in recs:
    k = (s['platform'], s.get('account_id'), str(s.get('ad_id')))
    first = 0 if k in firsts else 1; firsts.add(k)
    rows.append([s['platform'], s.get('account_name'), s.get('campaign_name'), s.get('adset_name'), s.get('ad_name') or '(no ad name)', str(s.get('ad_id')),
                 s['_kind'], s.get('effective_status') or s.get('ad_status'), 'Yes' if s['_on'] else 'No', s.get('slot'), s.get('url'),
                 s['verdict'], viol(s['violations']), st, why, target, after, vnow, first, agg(ad_before[k]), agg(ad_now[k])])
f1, l1 = table(ws, H, rows, 'EveryURL', [10, 24, 42, 42, 44, 22, 9, 16, 9, 28, 70, 12, 50, 18, 55, 70, 70, 11, 8, 11, 11],
               wrap_cols=(3, 4, 5, 11, 13, 15, 16, 17))
for col in ('Verdict before', 'Verdict now', 'Ad verdict before', 'Ad verdict now'):
    L = C[col]
    for val, fill in (('WRONG', RED), ('CORRECT', GREEN)):
        ws.conditional_formatting.add(f'{L}{f1+1}:{L}{l1}', FormulaRule(formula=[f'${L}{f1+1}="{val}"'], fill=fill))
L = C['Status now']
for val, fill in (('Fixed live', GREEN), ('Already correct', GREEN), ('Manual - you', AMBER), ('After reconnect', BLUE), ('Outcome unknown', BLUE), ('Flagged - not fixable', RED)):
    ws.conditional_formatting.add(f'{L}{f1+1}:{L}{l1}', FormulaRule(formula=[f'${L}{f1+1}="{val}"'], fill=fill))
LV = "'Every URL (live)'"
rng = lambda h: f"{LV}!${C[h]}${f1+1}:${C[h]}${l1}"

# ---------------- Manual - do yourself ----------------
wm = wb.create_sheet('Manual - do yourself')
WHERE = {'meta': 'Ads Manager > ad > Destination / Website URL (Meta creates a new creative when you save)',
         'x': 'X Ads > Creatives > Cards > edit card > Website URL',
         'tiktok': 'TikTok Ads Manager > ad > Destination URL (and URL parameters)',
         'snapchat': 'Snap Ads Manager > ad > Attachment > Website URL',
         'google': 'Google Ads > ad > Final URL', 'linkedin': 'Campaign Manager > ad > Landing page URL'}
mrows = []
for s, st, why, target, after, vnow in recs:
    if st != 'Manual - you': continue
    qa = dict(parse_qsl(urlsplit(s['url']).query)); qb = dict(parse_qsl(urlsplit(target).query))
    what = '\n'.join(f'{k}: {qa.get(k)} -> {qb.get(k)}' for k in sorted(set(qa) | set(qb)) if qa.get(k) != qb.get(k))
    m = manual.get((s['platform'], str(s.get('ad_id')), s.get('slot')))
    su = skipped_unit.get((s['platform'], str(s.get('ad_id'))))
    shared = m['where'].split(' | ')[0] if m and m['where'].startswith('shared') else (f"shared {su['unit']}" if su and why.startswith('Shared') else '')
    mrows.append([s['platform'], s.get('account_name'), s.get('campaign_name'), s.get('adset_name'), s.get('ad_name') or '(no ad name)', str(s.get('ad_id')),
                  s.get('effective_status') or s.get('ad_status'), 'Yes' if s['_on'] else 'No', why, s.get('slot') + (f' ({shared})' if shared else ''),
                  WHERE.get(s['platform'], ''), s['url'], target, what, ''])
mrows.sort(key=lambda r: (r[8][:20], r[0], str(r[1]), str(r[2]), str(r[4])))
fm, lm = table(wm, ['Platform', 'Account', 'Campaign', 'Ad set', 'Ad name', 'Ad ID', 'Delivery status', 'Switched on', 'Why it is manual', 'URL slot',
                    'Where to edit', 'URL now', 'Set the URL to (exact)', 'What changes', 'Done? (tick)'],
               mrows, 'Manual', [10, 22, 38, 38, 44, 22, 14, 9, 45, 30, 34, 70, 70, 50, 10], wrap_cols=(3, 4, 5, 9, 10, 11, 12, 13, 14))

# ---------------- After reconnect ----------------
wr = wb.create_sheet('After reconnect')
rrows = []
for s, st, why, target, after, vnow in recs:
    if st not in ('After reconnect', 'Outcome unknown'): continue
    fn = batch_of.get((s['platform'], str(s.get('ad_id')), s.get('slot')), '')
    if after: fn = 'second edit (was ' + fn + ')'
    rrows.append([s['platform'], s.get('account_name'), s.get('campaign_name'), s.get('adset_name'), s.get('ad_name') or '(no ad name)', str(s.get('ad_id')),
                  s.get('effective_status') or s.get('ad_status'), 'Yes' if s['_on'] else 'No', st, why, s.get('slot'), after or s['url'], target, fn])
rrows.sort(key=lambda r: (r[8], r[0], str(r[1]), str(r[2]), str(r[4])))
table(wr, ['Platform', 'Account', 'Campaign', 'Ad set', 'Ad name', 'Ad ID', 'Delivery status', 'Switched on', 'Status', 'Note', 'URL slot',
           'URL now', 'Will be set to (exact)', 'Batch'],
      rrows, 'Reconnect', [10, 22, 38, 38, 44, 22, 14, 9, 16, 50, 30, 70, 70, 16], wrap_cols=(3, 4, 5, 10, 12, 13))

# ---------------- Changes applied ----------------
wa = wb.create_sheet('Changes applied')
arows = []
byk = {}
for k, a in fixed.items(): byk[(a['_file'], a['unit'], a['slot'])] = a
for a in sorted(byk.values(), key=lambda a: (a['_run'], a['platform'], a['unit'])):
    qa = dict(parse_qsl(urlsplit(a['before']).query)); qb = dict(parse_qsl(urlsplit(a['after_live']).query))
    what = '\n'.join(f'{k}: {qa.get(k)} -> {qb.get(k)}' for k in sorted(set(qa) | set(qb)) if qa.get(k) != qb.get(k))
    arows.append([a['_run'], a['platform'], a.get('account_id'), '\n'.join(a.get('ad_names') or []), '\n'.join(map(str, a.get('ad_ids') or [])), a['slot'],
                  a['before'], a['after_live'], 'Yes' if a['after_live'] == a['after_expected'] else 'NO',
                  'Yes' if a['_run'].startswith('Generator') or all(TGT.get((a['platform'], str(x), a['slot'])) == a['after_live'] for x in a['ad_ids']) else 'NO - second edit needed', what, a.get('new_entity_ids') or '', (a.get('detail') or '')[:400]])
table(wa, ['Run', 'Platform', 'Account ID', 'Ad name', 'Ad ID', 'URL slot', 'Before', 'After (read back live)', 'Matches planned URL', 'Matches final rules', 'What changed', 'New entity (e.g. creative)', 'Platform response'],
      arows, 'Applied', [16, 10, 20, 44, 22, 28, 70, 70, 10, 14, 50, 22, 50], wrap_cols=(4, 5, 7, 8, 11, 13))

# ---------------- Flagged - not fixable ----------------
wf = wb.create_sheet('Flagged - not fixable')
frows = []
for s, st, why, target, after, vnow in recs:
    if st != 'Flagged - not fixable': continue
    frows.append([s['platform'], s.get('account_name'), s.get('campaign_name'), s.get('adset_name'), s.get('ad_name') or '(no ad name)', str(s.get('ad_id')),
                  s['_kind'], s.get('effective_status') or s.get('ad_status'), 'Yes' if s['_on'] else 'No', why, s['url'], viol(s['violations'])])
frows.sort(key=lambda r: (r[0], str(r[1]), r[9][:40], str(r[2])))
table(wf, ['Platform', 'Account', 'Campaign', 'Ad set', 'Ad name', 'Ad ID', 'Kind', 'Delivery status', 'Switched on', 'Why it cannot be fixed from the generator',
           'Live URL', 'What is wrong'], frows, 'Flagged', [10, 22, 38, 38, 44, 22, 9, 14, 9, 50, 70, 50], wrap_cols=(3, 4, 5, 10, 11, 12))

# ---------------- 7-Oct sheet audit ----------------
ws3 = wb.create_sheet('7-Oct sheet audit')
r3 = []
for a in SHEET:
    if not a.get('final_url'): continue
    codes = {e[0] for e in a['errors']}
    claim = a.get('outcome') or ''
    if a.get('rollout_result'): claim = (claim + ' | ' if claim else '') + 'Rollout: ' + str(a['rollout_result'])
    r3.append([a.get('platform'), a.get('campaign'), a.get('adset'), a.get('ad_name_after') or a.get('ad_name'), a['ad_id'], a['final_url'],
               a['final_src'], 'Yes' if a['claude_wrote_final'] else 'No', claim, 'No' if a['errors'] else 'Yes',
               'Yes' if codes - {'CAMPAIGN_NO_OBJECTIVE'} else 'No', viol(a['errors'])])
f3, l3 = table(ws3, ['Platform', 'Campaign', 'Ad set', 'Ad name', 'Ad ID', 'Final URL in the 7-Oct sheet', 'Tab it came from', 'Written by Claude',
                     'What the sheet claimed', 'Compliant with nomenclature', 'Breaks format rules (beyond objective)', 'What is wrong'],
               r3, 'SheetAudit', [12, 40, 40, 44, 22, 70, 22, 10, 34, 12, 14, 60], wrap_cols=(2, 3, 4, 6, 9, 12))
AU = "'7-Oct sheet audit'"
ar = lambda c: f"{AU}!${c}${f3+1}:${c}${l3}"

# ---------------- Rules ----------------
ws5 = wb.create_sheet('Rules used')
rules = [
    ('utm_source', 'Mappings col A list; must belong to the platform', 'Schema row 6'),
    ('utm_medium', 'Mappings col B list', 'Schema row 7'),
    ('utm_campaign', '{Campaign Unique ID}_{campaign name}_{mon}{yy}_{objective}, lower-case', 'Schema row 8 / Output formula'),
    ('utm_term', '{Adset Unique ID}_{ad set name}_{ad set language}; language from the AD SET', 'Schema row 9 / Output formula'),
    ('utm_content', '{Ad Unique ID}_{ad name}_{ad format}_{ad language}', 'Schema row 10 / Output formula'),
    ('Unique IDs', 'Only the 8-character IDs in Campaign_Code_Generator. Nothing invented (no ad-id-AD1001, adset-id-XXXXXX, RSA1, {adgroupid} macros).', 'Campaign_Code_Generator'),
    ('Names without a generator code', 'Legacy campaigns: the code is dropped where the generator has no code for that name; the name, language, format and objective are kept from the ad / ad set / campaign names. Where the generator has the name, its code is used.', 'User decision 11-Oct ("drop the code if it doesnt exist")'),
    ('Manazel', 'Generator codes used where the generator has them (e.g. 7ebeb0a8 manazel-alhadawi-2026); otherwise the code is dropped.', 'User decision 11-Oct'),
    ('Month (mon yy)', 'For legacy and month-clash ads: the month the ad first served (first day with impressions, account time zone). Ads that never served are flagged.', 'User decision 11-Oct'),
    ('Objective', 'Added to every utm_campaign: AWR->awr, CONV/LEAD->conv, TRF->cond, awareness/consideration/conversion words', 'User decision 11-Oct'),
    ('utm_id', 'Not added; removed where present', 'User decision 11-Oct'),
    ('kva / kvb / kvd', 'Programmatic only; on social or search the ad is flagged, not changed', 'User decision 11-Oct'),
    ('Search on social', 'Flagged, not changed', 'Your example'),
    ('Ritz Google search', 'utm_content = 78f964bb_rcsr-search_search_{lang}', 'User decision 11-Oct (Campaign_Code_Generator)'),
    ('Shared creatives / cards', 'Where one creative serves ads that need different URLs: listed for you to split by hand', 'User decision 11-Oct'),
    ('Sitelinks', 'No generator creative code exists for sitelinks: counted and flagged, not changed', 'No assumption'),
    ('Landing page', 'Kept as is; only the UTM query is rebuilt; click ids (fbclid etc.) removed', 'No change to user journey'),
    ('Scope', 'Every ad in every Meta, TikTok, Snapchat, LinkedIn, X and Google Ads account whose name contains "Diriyah" or "Wadi", pulled live 11-Oct-2026 via Whathead', 'User instruction 11-Oct'),
    ('Fixed live', 'Counted as fixed only when the URL was read back from the platform after the write and matched exactly', 'Verification'),
]
table(ws5, ['Rule', 'What is required', 'Source'], rules, 'Rules', [26, 110, 40], wrap_cols=(2,))

# ---------------- Summary ----------------
w0 = wb.create_sheet('Summary', 0)
w0['A1'] = 'Diriyah / Wadi Safar UTM verification vs Campaign Nomenclature Generator'; w0['A1'].font = TITLE
w0['A2'] = ('Live pull 11-Oct-2026 of every ad in every ad account whose name contains "Diriyah" or "Wadi" (Meta, TikTok, Snapchat, LinkedIn, X, Google Ads). '
            '"Now" = after the edits made so far, each read back from the platform after the write. All figures are formulas over the tabs in this file.')
w0['A2'].font = NOTE; w0['A2'].alignment = WRAP; w0.merge_cells('A2:N2'); w0.row_dimensions[2].height = 30
w0['A4'] = 'Ads and URL slots: before vs now (Google sitelinks shown separately)'; w0['A4'].font = H2
hdr = ['Platform', 'Ads with a URL', 'Ads wrong before', '% ads wrong before', 'Ads wrong now', '% ads wrong now', 'URL slots', 'Slots wrong before',
       '% slots wrong before', 'Fixed live', 'Manual - you', 'After reconnect / unknown', 'Flagged - not fixable', 'Slots wrong now', '% slots wrong now']
w0.append(hdr)
for c in w0[w0.max_row]: c.font = HDR; c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical='center')
w0.row_dimensions[w0.max_row].height = 42
plats = sorted({s['platform'] for s in rows_src})
def frow(label, extra):
    r = w0.max_row + 1
    base = f'{rng("Kind")},"Ad"' + extra
    first = f'{rng("First row of ad")},1'
    w0.append([label,
               f'=COUNTIFS({base},{first},{rng("Ad verdict before")},"<>NO URL")',
               f'=COUNTIFS({base},{first},{rng("Ad verdict before")},"WRONG")',
               f'=IF(B{r}=0,0,C{r}/B{r})',
               f'=COUNTIFS({base},{first},{rng("Ad verdict now")},"WRONG")',
               f'=IF(B{r}=0,0,E{r}/B{r})',
               f'=COUNTIFS({base},{rng("Verdict before")},"<>NO URL")',
               f'=COUNTIFS({base},{rng("Verdict before")},"WRONG")',
               f'=IF(G{r}=0,0,H{r}/G{r})',
               f'=COUNTIFS({base},{rng("Status now")},"Fixed live")',
               f'=COUNTIFS({base},{rng("Status now")},"Manual - you")',
               f'=COUNTIFS({base},{rng("Status now")},"After reconnect")+COUNTIFS({base},{rng("Status now")},"Outcome unknown")',
               f'=COUNTIFS({base},{rng("Status now")},"Flagged - not fixable")',
               f'=COUNTIFS({base},{rng("Verdict now")},"WRONG")',
               f'=IF(G{r}=0,0,N{r}/G{r})'])
    for c in w0[r]: c.font = BOLD if label.startswith(('ALL', 'Switched')) else BODY
    for L in 'DFIO': w0[f'{L}{r}'].number_format = '0.0%'
for p in plats: frow(p, f',{rng("Platform")},"{p}"')
frow('ALL ADS', '')
frow('Switched on now', f',{rng("Switched on")},"Yes"')
r = w0.max_row + 1
sl = f'{rng("Kind")},"Sitelink"'
w0.append(['Google sitelinks', '', '', '', '', '', f'=COUNTIFS({sl},{rng("Verdict before")},"<>NO URL")', f'=COUNTIFS({sl},{rng("Verdict before")},"WRONG")',
           f'=IF(G{r}=0,0,H{r}/G{r})', f'=COUNTIFS({sl},{rng("Status now")},"Fixed live")', '', '', f'=COUNTIFS({sl},{rng("Status now")},"Flagged - not fixable")',
           f'=COUNTIFS({sl},{rng("Verdict now")},"WRONG")', f'=IF(G{r}=0,0,N{r}/G{r})'])
for c in w0[r]: c.font = BODY
for L in 'IO': w0[f'{L}{r}'].number_format = '0.0%'
w0.append([]); w0.append(['An ad counts as wrong if any of its URL slots is wrong. "Ads with a URL" excludes ads with no readable URL (LinkedIn posts 403, Meta boosted posts, media-only tweets, TikTok ads without a landing page).'])
w0.cell(w0.max_row, 1).font = NOTE
w0.append([])
w0.append(['What is left']); w0.cell(w0.max_row, 1).font = H2
w0.append(['Item', 'URL slots', 'Where'])
for c in w0[w0.max_row]: c.font = HDR; c.fill = HFILL
for lbl, fml, where in [
        ('For you to change by hand (incl. shared creatives, TikTok ad-text trim)', f"=COUNTA('Manual - do yourself'!$A${fm+1}:$A${lm})", 'Tab "Manual - do yourself"'),
        ('Waiting for the Whathead reconnect (then applied and read back)', f'=COUNTIFS({rng("Status now")},"After reconnect")', 'Tab "After reconnect"'),
        ('Write sent as sign-in expired - re-read first', f'=COUNTIFS({rng("Status now")},"Outcome unknown")', 'Tab "After reconnect"'),
        ('Cannot be fixed from the generator without inventing (flagged)', f'=COUNTIFS({rng("Status now")},"Flagged - not fixable")', 'Tab "Flagged - not fixable"')]:
    w0.append([lbl, fml, where])
    for c in w0[w0.max_row]: c.font = BODY
w0.append([])
w0.append(['The 7-Oct "open_items_resolved" sheet']); w0.cell(w0.max_row, 1).font = H2
w0.append(['Measure', 'Count', '% of sheet URLs'])
for c in w0[w0.max_row]: c.font = HDR; c.fill = HFILL
b = w0.max_row
for i, (lbl, fml) in enumerate([
        ('URLs in the sheet (latest per ad)', f'=COUNTA({ar("F")})'),
        ('Not compliant with the nomenclature', f'=COUNTIFS({ar("J")},"No")'),
        ('Break format rules beyond the missing objective', f'=COUNTIFS({ar("K")},"Yes")'),
        ('Written or changed by Claude and not compliant', f'=COUNTIFS({ar("H")},"Yes",{ar("J")},"No")'),
        ('Marked "Correct"/"no change needed"/"Verified live" and not compliant', f'=COUNTIFS({ar("I")},"*Correct*",{ar("J")},"No")+COUNTIFS({ar("I")},"*Verified*",{ar("J")},"No")-COUNTIFS({ar("I")},"*Correct*",{ar("I")},"*Verified*",{ar("J")},"No")'),
        ('Contain an XXXX placeholder', f'=COUNTIFS({ar("L")},"*placeholder*")'),
        ('Use codes not in the generator (ad-id-/adset-id-)', f'=COUNTIFS({ar("L")},"*legacy*ad-id*")')]):
    rr = b + 1 + i
    w0.append([lbl, fml, f'=IF($B${b+1}=0,0,B{rr}/$B${b+1})'])
    for c in w0[rr]: c.font = BODY
    w0[f'C{rr}'].number_format = '0.0%'
for i, w in enumerate([62, 10, 10, 10, 10, 10, 10, 10, 10, 10, 10, 12, 12, 10, 10], 1): w0.column_dimensions[get_column_letter(i)].width = w
wb.save(OUT)

# ---------- console tally (for cross-checking the formulas) ----------
ad = [x for x in recs if x[0]['_kind'] == 'Ad']
url = [x for x in ad if x[0]['verdict'] != 'NO URL']
print('ad slots', len(url), 'wrong before', sum(x[0]['verdict'] == 'WRONG' for x in url), 'wrong now', sum(x[5] == 'WRONG' for x in url))
print('status', collections.Counter(x[1] for x in url))
fx = [x for x in url if x[1] == 'Fixed live']
print('fixed slots still failing checks', sum(x[5] == 'WRONG' for x in fx))
for x in fx:
    if x[5] == 'WRONG': print('   ', x[0]['platform'], x[0]['ad_id'], x[4], checks(x[0], x[4]))
ka = {}
for s, st, why, target, after, vnow in ad:
    k = (s['platform'], s.get('account_id'), str(s.get('ad_id'))); ka[k] = (agg(ad_before[k]), agg(ad_now[k]), s['_on'])
wa_ = [v for v in ka.values() if v[0] != 'NO URL']
print('ads with url', len(wa_), 'wrong before', sum(v[0] == 'WRONG' for v in wa_), 'wrong now', sum(v[1] == 'WRONG' for v in wa_))
on = [v for v in wa_ if v[2]]
print('switched-on ads', len(on), 'wrong before', sum(v[0] == 'WRONG' for v in on), 'wrong now', sum(v[1] == 'WRONG' for v in on))
print('manual rows', len(mrows), collections.Counter(r[8][:50] for r in mrows))
print('reconnect rows', len(rrows), collections.Counter((r[8], r[0]) for r in rrows))
print('applied rows', len(arows), collections.Counter((r[0], r[1]) for r in arows))
print('flagged rows', len(frows))
miss = [k for k in manual if k not in {(s['platform'], str(s.get('ad_id')), s.get('slot')) for s in rows_src}]
print('manual keys not matched to a live row', len(miss), miss[:5])
fk = [k for k in fixed if k not in {(s['platform'], str(s.get('ad_id')), s.get('slot')) for s in rows_src}]
print('fixed keys not matched to a live row', len(fk), fk[:5])
print('saved', OUT)
