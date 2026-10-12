"""Final verification workbook (12-Oct): before vs after %, every URL with its ad name, changes applied, what still needs the
user (not applied / manual), what was left out per URL, month check, questions, 7-Oct sheet audit, rules.
Usage: python3 scripts/build_final2.py <dir> <out.xlsx> [reread folder] [plan folder]
Reads: work/scored.json (11-Oct inventory = before), work/targets.json (final rules), work/live_now.json (keyed re-read = after),
       reread/out/new_*.jsonl (ads created since 11-Oct), apply*/out_*.jsonl (every write), round4/diff.json (pending plan), omit / legacy items."""
import sys, os, json, re, glob, collections
from urllib.parse import urlsplit, parse_qsl
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import FormulaRule
from derive import Deriver
from omit import check_omit

W, OUT = sys.argv[1], sys.argv[2]
REREAD = sys.argv[3] if len(sys.argv) > 3 else 'reread'   # folder of the latest keyed re-read (its out/new_*.jsonl = ads created since 11-Oct)
PLAN = sys.argv[4] if len(sys.argv) > 4 else 'round4'      # folder of the latest pending plan (diff.json)
G = glob.glob(f'{W}/in/e2e*.xlsx')[0]
N = Deriver(G, {}).N
S = json.load(open(f'{W}/work/scored.json'))
T = {(r['platform'], r['ad_id'], r['slot']): r for r in json.load(open(f'{W}/work/targets.json'))}
LIVE_FILE = os.environ.get('LIVE_FILE') or f'{W}/work/live_now.json'   # e.g. work/live_interim.json = re-read + later apply read-backs
LIVE = json.load(open(LIVE_FILE)) if os.path.exists(LIVE_FILE) else {}
DIFF = {(d['platform'], d['ad_id'], d['slot']): d for d in json.load(open(f'{W}/{PLAN}/diff.json'))}
PMAN = {}   # the plan's manual list: tweet-text links, shared sub-creatives needing different URLs
for m in json.load(open(f'{W}/{PLAN}/manual.json')):
    for aid in ([m['ad_id']] if m.get('ad_id') else m.get('ad_ids') or []):
        for sl_ in ([m['slot']] if m.get('slot') else [c_['slot'] for c_ in m.get('changes') or []]):
            PMAN[(m['platform'], str(aid), sl_)] = m.get('manual') if isinstance(m.get('manual'), str) else 'shared creative used by ads that need different URLs - split it by hand'
FS = json.load(open(f'{W}/work/first_serving_all.json'))
key = lambda s: (s['platform'], str(s.get('ad_id')), s.get('slot'))
OMIT = {key(o): o for o in json.load(open(f'{W}/work/omit_items.json'))}
LEG = {key(i): i for i in json.load(open(f'{W}/work/legacy_clash_items.json'))}
MONTHFIX = {(f['platform'], f['ad_id'], f['slot']): f for f in json.load(open(f'{W}/work/month_fixes.json'))}
norm = lambda u: (lambda p: (p.scheme, p.netloc.lower(), p.path, p.fragment, tuple(sorted(parse_qsl(p.query, keep_blank_values=True)))))(urlsplit((u or '').strip()))
same = lambda a, b: a is not None and b is not None and norm(a) == norm(b)
MAIN_SESSION = {('google', '812832821725'), ('google', '812832821731')}   # written from the main session on 12-Oct (see notes)

# ---------- every write ----------
runs = collections.defaultdict(list)
for f in sorted(glob.glob(f'{W}/apply*/out_*.jsonl')):
    day = '11-Oct' if re.search(r'/apply2?/', f) else '12-Oct'
    for line in open(f):
        if not line.strip(): continue
        try: a = json.loads(line)
        except Exception: continue
        a['_day'] = day; a['_file'] = os.path.relpath(f, W)
        for aid in a.get('ad_ids') or []:
            runs[(a['platform'], str(aid), a.get('slot'))].append(a)
        if a.get('unit_type') == 'sitelink_asset':      # one asset = every placement
            for k in T:
                if k[0] == 'google' and k[2].split(' final')[0] == 'sitelink asset ' + str(a['unit']): runs[k].append(a)

OFF = re.compile(r'PAUSED|DISABLE|INVALID|NOT_SERVING|REJECTED|DONE|ARCHIVED|REMOVED|ENDED|COMPLETED|DRAFT|HOLD', re.I)
def is_on(s, status=None):
    st = status if status is not None else (s.get('effective_status') or s.get('ad_status'))
    for k in ('campaign_status', 'adset_status'):
        v = s.get(k)
        if v is not None and not re.match(r'\s*(ACTIVE|ENABLE|ENABLED|SERVING)', str(v), re.I): return False
    return bool(st) and not OFF.search(str(st))

def wrong_before(s):
    return '\n'.join(m for _, m in (s.get('violations') or [])) or ('; '.join(s.get('derive_reasons') or []))

def status_of(s, t, live):
    """-> status, detail, changed_on"""
    k = key(s); tgt = (t or {}).get('target')
    rs = runs.get(k, [])
    last12 = [a for a in rs if a['_day'] == '12-Oct']
    if live is None: return 'Not re-read', 'the final re-read did not return this slot', ''
    if live.get('found') is False: return 'Gone', live.get('note') or 'ad / slot no longer exists', ''
    url = live.get('url')
    if tgt and same(url, tgt):
        if (s['platform'], str(s.get('ad_id'))) in MAIN_SESSION: return 'Fixed 12-Oct', 'written from the main session (agent stopped by the safety check)', '12-Oct'
        if any(a['result'] in ('Fixed',) for a in last12): return 'Fixed 12-Oct', next(a.get('detail') for a in last12 if a['result'] == 'Fixed')[:300], '12-Oct'
        if any(a['result'] == 'Fixed' for a in rs): return 'Fixed 11-Oct', 'written 11-Oct and still matches the final rules', '11-Oct'
        if any(a['result'] == 'AlreadyFixed' for a in last12): return 'Fixed 11-Oct', 'already correct when re-read before the 12-Oct write', '11-Oct'
        return 'Correct', 'live URL already matches the final rules', ''
    if t and t.get('manual'): return 'Manual - you', t['manual'], ''
    if k in PMAN: return 'Manual - you', PMAN[k], ''
    if not tgt:
        o = OMIT.get(k) or {}
        return 'Manual - you', (o.get('manual') or 'no single correct URL for this slot'), ''
    a5 = [a for a in rs if a['_file'].startswith(PLAN + '/')]
    if k in DIFF and not a5: return 'Being applied now', 'reviewed URL is in the write queue that was still running when this workbook was built', ''
    bad = [a for a in (a5 or last12) if a['result'] in ('Failed', 'Skipped', 'NotAttempted')]
    if bad:
        b = bad[-1]; d = (b.get('detail') or '')
        if b['result'] == 'NotAttempted' or re.search(r'safety check|sign-?in', d, re.I):
            return 'Not applied - needs your OK', 'planned and reviewed, not written: the session safety check stopped this live write', ''
        if re.search(r'Spark Pull|editable only in Ads Manager|archived', d, re.I):
            return 'Manual - you', 'the platform does not allow this edit through its API: ' + d[:300], ''
        return {'Failed': 'Failed', 'Skipped': 'Skipped'}[b['result']], d[:400], ''
    if any(a['result'] == 'Fixed' and same(url, a.get('after_live')) for a in rs):
        return 'Not applied - needs your OK', 'written earlier with a previous version of the rules; the final rules change it once more (second edit ready, not written: the session safety check stopped further live writes)', ''
    if k in DIFF or any(same(url, x) for x in [s.get('url')] + [a.get('after_live') for a in rs]):
        return 'Not applied - needs your OK', 'planned and reviewed, not written: the session safety check stopped further live writes - see "Not applied - needs your OK"', ''
    return 'Changed by someone else', 'live URL differs from both the 11-Oct URL and the planned one', ''

def left_out(k, t):
    o = OMIT.get(k)
    if t and (t.get('source') or '').startswith('omit') and o: return ', '.join(o.get('omitted') or [])
    return ''

def how(k, t):
    b = []
    i = LEG.get(k) or OMIT.get(k)
    if i: b += i.get('basis') or []
    if t: b += (t.get('consistency') or []) + ([t['month_fix']] if t.get('month_fix') else [])
    return '\n'.join(dict.fromkeys(str(x) for x in b if x))[:1500]

# ---------- one row per ad URL slot, one per sitelink asset ----------
recs, seen = [], set()
for s in S:
    if not s.get('url'): continue
    k = key(s); t = T.get(k)
    kind = 'Sitelink' if 'sitelink' in str(s.get('slot')).lower() else 'Ad'
    if kind == 'Sitelink':
        a = (s['account_id'], str(s.get('ad_id')))
        if a in seen: continue
        seen.add(a)
    live = LIVE.get('|'.join(k))
    st, why, day = status_of(s, t, live)
    url_now = (live or {}).get('url')
    vnow = 'GONE' if st == 'Gone' else ('CORRECT' if t and t.get('target') and same(url_now, t['target']) else 'WRONG')
    if st == 'Not re-read': vnow = 'NOT RE-READ'
    status_now = (live or {}).get('status_now') or s.get('effective_status') or s.get('ad_status')
    recs.append(dict(s=s, t=t, kind=kind, st=st, why=why, day=day, url_now=url_now, vnow=vnow, status_now=status_now,
                     on=is_on(s, (live or {}).get('status_now')), before='WRONG' if s['verdict'] != 'CORRECT' else 'CORRECT',
                     placements=sum(1 for x in S if x.get('url') and kind == 'Sitelink' and x['account_id'] == s['account_id'] and str(x['ad_id']) == str(s['ad_id']))))
recs.sort(key=lambda r: (r['s']['platform'], str(r['s'].get('account_name')), r['kind'], str(r['s'].get('campaign_name')), str(r['s'].get('adset_name')), str(r['s'].get('ad_name')), str(r['s'].get('ad_id')), str(r['s'].get('slot'))))
ad_b, ad_n = collections.defaultdict(list), collections.defaultdict(list)
for r in recs:
    if r['kind'] == 'Ad':
        ka = (r['s']['platform'], r['s'].get('account_id'), str(r['s'].get('ad_id')))
        ad_b[ka].append(r['before']); ad_n[ka].append(r['vnow'])
agg = lambda v: 'WRONG' if 'WRONG' in v else ('CORRECT' if all(x == 'CORRECT' for x in v) else v[0])

# ---------- styling helpers ----------
F = 'Arial'
HDR = Font(name=F, bold=True, color='FFFFFF'); HFILL = PatternFill('solid', fgColor='1F3864')
BODY = Font(name=F, size=10); BOLD = Font(name=F, size=10, bold=True)
TITLE = Font(name=F, size=14, bold=True); H2 = Font(name=F, size=12, bold=True); NOTE = Font(name=F, size=9, italic=True, color='595959')
WRAP = Alignment(wrap_text=True, vertical='top'); TOP = Alignment(vertical='top')
RED = PatternFill('solid', fgColor='F8CBAD'); GREEN = PatternFill('solid', fgColor='C6EFCE'); AMBER = PatternFill('solid', fgColor='FFE699'); BLUE = PatternFill('solid', fgColor='DDEBF7')
def table(ws, headers, rows, name, widths, wrap_cols=()):
    ws.append(headers); hr = ws.max_row
    for c in ws[hr]: c.font = HDR; c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical='center')
    for r in rows:
        ws.append([('' if v is None else v) for v in r])
        for i, c in enumerate(ws[ws.max_row], 1): c.font = BODY; c.alignment = WRAP if i in wrap_cols else TOP
    last = max(ws.max_row, hr + 1)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    t = Table(displayName=name, ref=f'A{hr}:{get_column_letter(len(headers))}{last}')
    t.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True)
    ws.add_table(t); ws.freeze_panes = ws.cell(hr + 1, 1)
    return hr, last
def colour(ws, col, first, last, pairs):
    for val, fill in pairs:
        ws.conditional_formatting.add(f'{col}{first}:{col}{last}', FormulaRule(formula=[f'${col}{first}="{val}"'], fill=fill))

wb = Workbook()
# ---------------- Every URL ----------------
ws = wb.active; ws.title = 'Every URL'
H = ['Platform', 'Account', 'Campaign', 'Ad set / ad group', 'Ad name', 'Ad ID', 'Kind', 'Delivery status now', 'Switched on', 'URL slot',
     'URL before fix (11-Oct)', 'Verdict before', 'What was wrong before', 'Correct URL (final rules)', 'Built by', 'Parts left out (12-Oct rule)',
     'Live URL now (re-read)', 'Verdict now', 'Status', 'Detail / next step', 'Fixed on', 'Sitelink placements', 'First row of ad', 'Ad verdict before', 'Ad verdict now']
C = {h: get_column_letter(i) for i, h in enumerate(H, 1)}
rows, firsts = [], set()
for r in recs:
    s, t = r['s'], r['t']
    ka = (s['platform'], s.get('account_id'), str(s.get('ad_id')))
    first = 1 if (r['kind'] == 'Ad' and ka not in firsts) else 0
    if r['kind'] == 'Ad': firsts.add(ka)
    rows.append([s['platform'], s.get('account_name'), s.get('campaign_name'), s.get('adset_name'), s.get('ad_name') or '(no ad name)', str(s.get('ad_id')),
                 r['kind'], r['status_now'], 'Yes' if r['on'] else 'No', s.get('slot'), s.get('url'), r['before'], wrong_before(s),
                 (t or {}).get('target') or '', (t or {}).get('source') or '', left_out(key(s), t), r['url_now'], r['vnow'], r['st'], r['why'], r['day'],
                 r['placements'] if r['kind'] == 'Sitelink' else '', first,
                 agg(ad_b[ka]) if r['kind'] == 'Ad' else r['before'], agg(ad_n[ka]) if r['kind'] == 'Ad' else r['vnow']])
f1, l1 = table(ws, H, rows, 'EveryURL', [9, 22, 40, 38, 44, 22, 8, 16, 8, 26, 70, 10, 40, 70, 10, 34, 70, 11, 20, 50, 9, 9, 7, 10, 10],
               wrap_cols=(3, 4, 5, 11, 13, 14, 16, 17, 20))
for col in ('Verdict before', 'Verdict now', 'Ad verdict before', 'Ad verdict now'):
    colour(ws, C[col], f1 + 1, l1, (('WRONG', RED), ('CORRECT', GREEN)))
colour(ws, C['Status'], f1 + 1, l1, (('Fixed 12-Oct', GREEN), ('Fixed 11-Oct', GREEN), ('Correct', GREEN), ('Manual - you', AMBER),
                                      ('Not applied - needs your OK', BLUE), ('Being applied now', BLUE), ('Failed', RED), ('Skipped', AMBER), ('Changed by someone else', AMBER)))
EU = "'Every URL'"
rng = lambda h: f"{EU}!${C[h]}${f1+1}:${C[h]}${l1}"

# ---------------- Changes applied ----------------
wa = wb.create_sheet('Changes applied')
ca = []
for f in sorted(glob.glob(f'{W}/apply*/out_*.jsonl')):
    day = '11-Oct' if re.search(r'/apply2?/', f) else '12-Oct'
    for line in open(f):
        if not line.strip(): continue
        try: a = json.loads(line)
        except Exception: continue
        if a['result'] not in ('Fixed', 'AlreadyFixed'): continue
        ok = same(a.get('after_live'), a.get('after_expected'))
        tg = None
        for aid in a.get('ad_ids') or []:
            tt = T.get((a['platform'], str(aid), a.get('slot')))
            if tt: tg = tt.get('target'); break
        ca.append([day, a['platform'], a.get('account_id'), '; '.join(str(x) for x in (a.get('ad_names') or []) if x) or '(no ad name)',
                   ', '.join(str(x) for x in a.get('ad_ids') or []), str(a.get('unit')), a.get('slot'), a.get('before'), a.get('after_live'),
                   'Yes' if ok else 'No', ('Yes' if same(a.get('after_live'), tg) else 'No - superseded, see Every URL') if tg else '',
                   a['result'], a.get('new_entity_ids') or '', (a.get('detail') or '')[:400]])
for (p, aid) in sorted(MAIN_SESSION):
    s_ = next((x for x in S if x['platform'] == p and str(x.get('ad_id')) == aid), None)
    lv = LIVE.get(f"{p}|{aid}|{s_['slot']}") if s_ else None
    tg = (T.get((p, aid, s_['slot'])) or {}).get('target') if s_ else None
    ca.append(['12-Oct', p, s_['account_id'] if s_ else '', '(no ad name)', aid, aid, s_['slot'] if s_ else '', s_['url'] if s_ else '', (lv or {}).get('url') or '',
               'Yes' if lv and same(lv.get('url'), tg) else '', 'Yes' if lv and same(lv.get('url'), tg) else '', 'Fixed',
               '', 'written from the main session after the agent was stopped by the safety check (dry-run, MutateAds final_urls, GAQL read-back byte-equal, status ENABLED unchanged)'])
table(wa, ['Day', 'Platform', 'Account ID', 'Ad name(s)', 'Ad ID(s)', 'Edited entity', 'URL slot', 'Before', 'After (read back live)', 'Read-back = planned',
           'Matches final rules', 'Result', 'New entity (e.g. creative)', 'Detail'], ca, 'ChangesApplied',
      [8, 9, 14, 44, 22, 22, 26, 70, 70, 10, 14, 11, 16, 60], wrap_cols=(4, 8, 9, 14))

# ---------------- Not applied - needs your OK ----------------
wn = wb.create_sheet('Not applied - needs your OK')
na = [[r['s']['platform'], r['s'].get('account_name'), r['s'].get('campaign_name'), r['s'].get('adset_name'), r['s'].get('ad_name') or '(no ad name)',
       str(r['s'].get('ad_id')), r['kind'], r['status_now'], 'Yes' if r['on'] else 'No', r['s'].get('slot'), r['url_now'], (r['t'] or {}).get('target'), r['why']]
      for r in recs if r['st'] in ('Not applied - needs your OK', 'Being applied now')]
table(wn, ['Platform', 'Account', 'Campaign', 'Ad set', 'Ad name', 'Ad ID', 'Kind', 'Delivery status', 'Switched on', 'URL slot', 'URL now', 'Correct URL (ready, reviewed)', 'Why not applied'],
      na, 'NotApplied', [9, 22, 40, 36, 44, 22, 8, 16, 8, 26, 70, 70, 50], wrap_cols=(3, 4, 5, 11, 12, 13))

# ---------------- Manual - do yourself ----------------
wm = wb.create_sheet('Manual - do yourself')
WHERE = {'meta': 'Ads Manager > ad > Edit > Website URL', 'snapchat': 'Snap Ads Manager > ad > Attachment > Website URL', 'tiktok': 'TikTok Ads Manager > ad > Destination URL',
         'x': 'X Ads > creative (card) > Website URL', 'google': 'Google Ads > Assets > Sitelinks (split the asset per campaign type)', 'linkedin': 'Campaign Manager > ad > Landing page'}
ma = [[r['s']['platform'], r['s'].get('account_name'), r['s'].get('campaign_name'), r['s'].get('adset_name'), r['s'].get('ad_name') or '(no ad name)', str(r['s'].get('ad_id')),
       r['kind'], r['status_now'], 'Yes' if r['on'] else 'No', r['st'], r['why'], r['s'].get('slot'), WHERE.get(r['s']['platform']), r['url_now'], (r['t'] or {}).get('target') or '']
      for r in recs if r['st'] in ('Manual - you', 'Failed', 'Skipped', 'Changed by someone else')]
# landing pages whose language contradicts the ad (target keeps the ad's language; the page itself is the user's call)
for r in recs:
    s = r['s']; tg = (r['t'] or {}).get('target')
    if not tg or r['kind'] != 'Ad': continue
    lp = re.search(r'/(en|ar)(/|$)', urlsplit(tg).path.lower()); q = dict(parse_qsl(urlsplit(tg).query))
    cl = (q.get('utm_content', '').split('_') or [''])[-1]; tl = (q.get('utm_term', '').split('_') or [''])[-1]
    cl = cl if cl in ('en', 'ar') else (tl if tl in ('en', 'ar') else None)   # Google ads have no ad name: fall back to the ad group language
    if lp and cl and cl != lp.group(1):
        ma.append([s['platform'], s.get('account_name'), s.get('campaign_name'), s.get('adset_name'), s.get('ad_name') or '(no ad name)', str(s.get('ad_id')), r['kind'], r['status_now'],
                   'Yes' if r['on'] else 'No', 'Check landing page', f'the ad / ad group is {cl.upper()} but its landing page is /{lp.group(1)}/ - the UTM follows the ad; change the page if it is wrong',
                   s.get('slot'), WHERE.get(s['platform']), r['url_now'], tg])
table(wm, ['Platform', 'Account', 'Campaign', 'Ad set', 'Ad name', 'Ad ID', 'Kind', 'Delivery status', 'Switched on', 'Status', 'Why it is yours', 'URL slot', 'Where to edit',
           'URL now', 'Correct URL (when one exists)'], ma, 'Manual', [9, 22, 40, 36, 44, 22, 8, 16, 8, 16, 55, 26, 34, 70, 70], wrap_cols=(3, 4, 5, 11, 14, 15))

# ---------------- What was left out ----------------
wl = wb.create_sheet('What was left out')
lo = []
for r in recs:
    k = key(r['s']); t = r['t']
    if not t: continue
    lo.append([r['s']['platform'], r['s'].get('ad_name') or '(no ad name)', str(r['s'].get('ad_id')), r['kind'], t.get('source'), t.get('target'), left_out(k, t) or '(nothing - every part sourced)', how(k, t)])
table(wl, ['Platform', 'Ad name', 'Ad ID', 'Kind', 'Built by', 'Correct URL (final rules)', 'Parts left out', 'How each part was sourced'], lo, 'LeftOut',
      [9, 44, 22, 8, 10, 70, 40, 90], wrap_cols=(2, 6, 7, 8))

# ---------------- Month check ----------------
wmo = wb.create_sheet('Month check')
mo = []
for r in recs:
    s = r['s']
    if r['kind'] != 'Ad': continue
    fs = FS.get(f"{s['platform']}|{s.get('ad_id')}")
    tg = (r['t'] or {}).get('target') or ''
    m = re.search(r'(?:^|_)((?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)\d\d)(?:_|$)', dict(parse_qsl(urlsplit(tg).query)).get('utm_campaign', '')) if tg else None
    mf = MONTHFIX.get(key(s))
    mo.append([s['platform'], s.get('ad_name') or '(no ad name)', str(s.get('ad_id')), fs or '', m.group(1) if m else '(left out)', (mf or {}).get('why') or ''])
table(wmo, ['Platform', 'Ad name', 'Ad ID', 'First day with impressions (account time zone)', 'Month in utm_campaign', 'Correction made on 12-Oct'], mo, 'MonthCheck',
      [9, 50, 22, 18, 14, 70], wrap_cols=(2, 6))

# ---------------- New since 11-Oct ----------------
wnw = wb.create_sheet('New ads since 11-Oct')
nw = []
for f in sorted(glob.glob(f'{W}/{REREAD}/out/new_*.jsonl')):
    for line in open(f):
        if not line.strip(): continue
        try: x = json.loads(line)
        except Exception: continue
        e = check_omit(N, x['url']) if x.get('url') else ['no URL']
        nw.append([x.get('platform'), x.get('account_name'), x.get('campaign_name'), x.get('adset_name'), x.get('ad_name'), str(x.get('ad_id')), x.get('effective_status') or x.get('ad_status'),
                   x.get('slot'), x.get('url'), 'passes the format checks' if not e else '; '.join(e)])
table(wnw, ['Platform', 'Account', 'Campaign', 'Ad set', 'Ad name', 'Ad ID', 'Delivery status', 'URL slot', 'URL now', 'Format check (not counted in the %)'], nw or [['(none found)'] + [''] * 9],
      'NewAds', [9, 22, 40, 36, 44, 22, 16, 26, 70, 50], wrap_cols=(3, 4, 5, 9, 10))

# ---------------- 7-Oct sheet audit (carried over) ----------------
old = glob.glob(f'{W}/out/Diriyah_UTM_Verification_2026-10-11_final.xlsx')
if old:
    ow = openpyxl.load_workbook(old[0])['7-Oct sheet audit']
    w7 = wb.create_sheet('7-Oct sheet audit')
    vals = [list(r) for r in ow.iter_rows(values_only=True)]
    table(w7, [str(x) for x in vals[0]], vals[1:], 'Audit7Oct', [9, 40, 36, 44, 22, 70, 22, 9, 40, 12, 12, 60], wrap_cols=(2, 3, 4, 6, 9, 12))

# ---------------- Questions for you ----------------
wq = wb.create_sheet('Questions for you')
QS = [
    ('Live writes still to do', 'You confirmed the remaining writes on 12-Oct. They were running when this workbook was built ("Being applied now"). The session safety check again stopped some Google and TikTok writes even after your OK; those are in "Not applied - needs your OK". To let them through, add a permission rule allowing the Whathead Google Ads / TikTok write actions, then they are applied with the same re-read / dry-run / read-back routine.', 'In progress'),
    ('utm_medium', 'Old URLs mixed cpc / cpm / paid-social / paid-search, even inside one campaign and on Google Search. The planned URLs use the generator\'s own medium for each source: google -> paid-search; meta / tiktok / snapchat / x / linkedin -> paid-social. Keep, or keep the old medium where it is a valid Mappings value?', 'Default applied: generator medium'),
    ('Demand Gen (source gadg)', 'The generator has no row for gadg, so no medium can be sourced; the planned URLs leave utm_medium out (12-Oct omit rule). Pick paid-display or paid-video if you want one.', 'Left out'),
    ('Template campaign name "visitor-diriyah-coe"', 'The agency template token (cam-id_mindshare_diriyah-company_<asset>_tactical_visitor-diriyah-coe_...) was pasted on Bujairi, Zallal, Bab Samhan and Shopping 2026 campaigns. Only its Campaign Name field is used (visitor-diriyah-coe), as the schema says; that merges those campaigns under one name in reports. Alternatives: leave the campaign name out, or keep the unit + asset (diriyah-company-bujairi-visitor-diriyah-coe).', 'Default applied: name field only'),
    ('Agency prefix "hs-"', 'Legacy H&S campaign tokens keep their hs- prefix (hs-cityscape, hs-wintercampaign-zallal). Strip it?', 'Kept'),
    ('utm_source', 'Where the campaign name carries its channel, the source follows it (instagram -> meta on _Meta_ campaigns, twitter -> x on _X_ campaigns), as the generator does.', 'Default applied'),
    ('Ritz rule', 'The 78f964bb_rcsr-search_search_{lang} content is applied to the Ritz-Carlton Signature (2026) search ads only, not to RitzCarlton_Residences 2024. Correct?', 'Signature only'),
    ('Landing pages', 'Some Arabic ads land on /en/ pages (and the reverse). The UTM language follows the ad; the page itself is listed in "Manual - do yourself" for you to check.', 'Your call'),
]
table(wq, ['Topic', 'Question / what was done', 'Status'], QS, 'Questions', [30, 120, 26], wrap_cols=(2,))

# ---------------- Rules used ----------------
wr = wb.create_sheet('Rules used')
RULES = [
    ('utm_campaign', 'uid_campaign-name_monYY_objective (all lower case)', 'Generator Schema / Output formula'),
    ('utm_term', 'uid_adset-name_adset-language', 'Generator Schema / Output formula'),
    ('utm_content', 'uid_ad-name_format_ad-language', 'Generator Schema / Output formula'),
    ('utm_source / utm_medium', 'From the Mappings lists; source follows the campaign name channel when it has one (_Meta_ -> meta, _X_ -> x)', 'Mappings; v5 review'),
    ('No utm_id', 'Removed wherever it appeared; click ids (fbclid, gclid) dropped', 'User 11-Oct'),
    ('Objective', 'Always added when known: _MF_TRF_ -> cond, AWR -> awr, CONV/LEAD -> conv, or the definition / sub-objective words (traffic -> cond, reach -> awr, leads/sales -> conv)', 'User 11-Oct; generator Input rows'),
    ('Uids', 'Only Campaign_Code_Generator 8-hex codes, only for exactly that name; otherwise no code. Manazel: generator code if it has one', 'User 11-Oct'),
    ('Month', 'Month the ad FIRST served (first day with impressions, account time zone), confirmed with daily data. Never served -> left out. Sitelinks -> left out', 'User 11-Oct ("be extremely accurate")'),
    ('kva / kvb / kvd', 'Programmatic-only: on social / search the ad token is unknown and left out', 'User 11-Oct'),
    ('Ritz Google search', '78f964bb_rcsr-search_search_{lang}', 'User 11-Oct'),
    ('12-Oct omit rule', '"for unknown figures like ad tokens write the utm without them" - a part that cannot be sourced is left out, never invented; a parameter with no known part is dropped', 'User 12-Oct'),
    ('Names (legacy ads)', 'From the generator, from names sibling ads in the same live campaign / ad set already use, or from the ad\'s own old UTM token; one name per live campaign / ad set; tokens that name another campaign or another cut are not used', 'v5 review (12-Oct)'),
    ('Format', 'Video slot -> video; carousel has no Mappings format and is written as display', 'Mappings col T; v5 review'),
    ('Shared creatives', 'A creative / card / sitelink used by ads that need different URLs is on the manual list (split it by hand)', 'User 11-Oct'),
    ('utm_medium (v6)', "Legacy / rebuilt URLs use the generator's medium for the source (google -> paid-search, social -> paid-social); Demand Gen (no generator row) left out", 'Generator Input rows; see Questions'),
    ('Old generator-layout tokens', 'agency_unit_asset_type_NAME_... : only the NAME field is a campaign name', 'Generator Schema; review 12-Oct'),
    ('Every write', 're-read live -> skip if it changed -> dry-run -> write -> read back and compare byte for byte; nothing deleted', 'Handoff'),
]
table(wr, ['Rule', 'What is required', 'Source'], RULES, 'Rules', [24, 110, 34], wrap_cols=(2,))

# ---------------- Summary ----------------
sm = wb.create_sheet('Summary', 0)
sm['A1'] = 'Diriyah / Wadi Safar UTM fix - before vs after' + (' (INTERIM - writes still running)' if os.environ.get('INTERIM') else ''); sm['A1'].font = TITLE
sm['A2'] = ('Every ad and every final-URL slot in every ad account whose name contains "Diriyah" or "Wadi" (Meta, TikTok, Snapchat, X, Google Ads, LinkedIn). '
            'Before = live pull 11-Oct-2026. After = keyed live re-read of the same slots on 12-Oct-2026, compared with the URL the final rules give. Ad name is on every row of "Every URL".')
sm['A2'].font = NOTE; sm['A2'].alignment = WRAP; sm.merge_cells('A2:N2'); sm.row_dimensions[2].height = 42
sm['A4'] = 'Ad URL slots and ads'; sm['A4'].font = H2
hdr = ['Platform', 'URL slots', 'Slots wrong before', '% wrong before', 'Slots wrong now', '% wrong now', 'Fixed 12-Oct', 'Fixed 11-Oct', 'Not applied - needs your OK',
       'Manual - you', 'Failed / skipped', 'Ads with a URL', 'Ads wrong before', '% ads wrong before', 'Ads wrong now', '% ads wrong now', 'Being applied now']
for i, h in enumerate(hdr, 1):
    c = sm.cell(5, i, h); c.font = HDR; c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical='center')
plats = ['meta', 'tiktok', 'snapchat', 'x', 'google', 'linkedin']
row = 6
P, K, VB, VN, ST, FR = rng('Platform'), rng('Kind'), rng('Verdict before'), rng('Verdict now'), rng('Status'), rng('First row of ad')
AB, AN = rng('Ad verdict before'), rng('Ad verdict now')
for p in plats + ['ALL']:
    pc = f'"{p}"' if p != 'ALL' else '"*"'
    sm.cell(row, 1, p if p != 'ALL' else 'All platforms').font = BOLD if p == 'ALL' else BODY
    f = lambda extra: f'=COUNTIFS({P},{pc},{K},"Ad"{extra})'
    sm.cell(row, 2, f('')); sm.cell(row, 3, f(f',{VB},"WRONG"')); sm.cell(row, 4, f'=IF(B{row}=0,0,C{row}/B{row})')
    sm.cell(row, 5, f(f',{VN},"WRONG"')); sm.cell(row, 6, f'=IF(B{row}=0,0,E{row}/B{row})')
    sm.cell(row, 7, f(f',{ST},"Fixed 12-Oct"')); sm.cell(row, 8, f(f',{ST},"Fixed 11-Oct"')); sm.cell(row, 9, f(f',{ST},"Not applied - needs your OK"'))
    sm.cell(row, 10, f(f',{ST},"Manual - you"')); sm.cell(row, 11, f'=' + f(f',{ST},"Failed"')[1:] + '+' + f(f',{ST},"Skipped"')[1:])
    sm.cell(row, 12, f(f',{FR},1')); sm.cell(row, 13, f(f',{FR},1,{AB},"WRONG"')); sm.cell(row, 14, f'=IF(L{row}=0,0,M{row}/L{row})')
    sm.cell(row, 15, f(f',{FR},1,{AN},"WRONG"')); sm.cell(row, 16, f'=IF(L{row}=0,0,O{row}/L{row})'); sm.cell(row, 17, f(f',{ST},"Being applied now"'))
    for c in (4, 6, 14, 16): sm.cell(row, c).number_format = '0.0%'
    for c in range(2, 18): sm.cell(row, c).font = BOLD if p == 'ALL' else BODY
    row += 1
row += 1
sm.cell(row, 1, 'Switched-on ads (delivering now)').font = H2; row += 1
for i, h in enumerate(['Scope', 'Ads', 'Wrong before', '% before', 'Wrong now', '% now'], 1):
    c = sm.cell(row, i, h); c.font = HDR; c.fill = HFILL
row += 1
ON = rng('Switched on')
sm.cell(row, 1, 'Switched-on ads'); sm.cell(row, 2, f'=COUNTIFS({K},"Ad",{FR},1,{ON},"Yes")'); sm.cell(row, 3, f'=COUNTIFS({K},"Ad",{FR},1,{ON},"Yes",{AB},"WRONG")')
sm.cell(row, 4, f'=IF(B{row}=0,0,C{row}/B{row})'); sm.cell(row, 5, f'=COUNTIFS({K},"Ad",{FR},1,{ON},"Yes",{AN},"WRONG")'); sm.cell(row, 6, f'=IF(B{row}=0,0,E{row}/B{row})')
for c in (4, 6): sm.cell(row, c).number_format = '0.0%'
row += 2
sm.cell(row, 1, 'Google sitelinks (one row per asset)').font = H2; row += 1
for i, h in enumerate(['Scope', 'Assets', 'Wrong before', '% before', 'Wrong now', '% now', 'Placements'], 1):
    c = sm.cell(row, i, h); c.font = HDR; c.fill = HFILL
row += 1
PL = rng('Sitelink placements')
sm.cell(row, 1, 'Sitelink assets'); sm.cell(row, 2, f'=COUNTIFS({K},"Sitelink")'); sm.cell(row, 3, f'=COUNTIFS({K},"Sitelink",{VB},"WRONG")'); sm.cell(row, 4, f'=IF(B{row}=0,0,C{row}/B{row})')
sm.cell(row, 5, f'=COUNTIFS({K},"Sitelink",{VN},"WRONG")'); sm.cell(row, 6, f'=IF(B{row}=0,0,E{row}/B{row})'); sm.cell(row, 7, f'=SUMIFS({PL},{K},"Sitelink")')
for c in (4, 6): sm.cell(row, c).number_format = '0.0%'
row += 2
sm.cell(row, 1, 'How to read it').font = H2; row += 1
notes = [
    '"Wrong now" counts every slot whose live URL is not the URL the final rules give - that includes the ones still waiting for your OK and the manual ones.',
    '"Not applied - needs your OK": planned and reviewed, but the session safety check stopped the live write. "Being applied now": in the write queue that was still running when this workbook was built. The ready URL is in "Not applied - needs your OK".',
    '"Manual - you": shared creatives / cards / sitelinks that would need two different URLs (split them by hand), plus anything a platform refused.',
    'Every write was re-read before and after; "Changes applied" has the before / after read back for each one.',
    'Ads created after the 11-Oct pull are listed in "New ads since 11-Oct" and are not counted in the percentages.',
]
for n in notes:
    sm.cell(row, 1, '- ' + n).font = BODY; sm.merge_cells(start_row=row, start_column=1, end_row=row, end_column=14); sm.cell(row, 1).alignment = WRAP; sm.row_dimensions[row].height = 28; row += 1
for i, w in enumerate([22, 10, 12, 10, 12, 10, 11, 11, 14, 11, 11, 11, 12, 10, 12, 10, 12], 1): sm.column_dimensions[get_column_letter(i)].width = w
sm.freeze_panes = 'A6'
wb.save(OUT)
# python-side numbers (for the summary check)
ad = [r for r in recs if r['kind'] == 'Ad']; sl = [r for r in recs if r['kind'] == 'Sitelink']
c = collections.Counter(r['st'] for r in ad)
print('ad slots', len(ad), 'wrong before', sum(r['before'] == 'WRONG' for r in ad), 'wrong now', sum(r['vnow'] == 'WRONG' for r in ad), dict(c))
ads = {(r['s']['platform'], r['s'].get('account_id'), str(r['s'].get('ad_id'))) for r in ad}
print('ads', len(ads), 'wrong before', sum(agg(ad_b[a]) == 'WRONG' for a in ads), 'wrong now', sum(agg(ad_n[a]) == 'WRONG' for a in ads))
print('sitelinks', len(sl), 'wrong now', sum(r['vnow'] == 'WRONG' for r in sl))
