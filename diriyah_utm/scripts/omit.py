"""User rule 12-Oct: "for unknown figures like ad tokens write the utm without them".
Where a part of the generator schema cannot be sourced (no generator code / no usable name / never served / kva on social or search),
the part is LEFT OUT instead of being invented or the URL being flagged:
  utm_campaign = [uid_][campaign name_][monYY_][objective]
  utm_term     = [uid_][ad set name_][ad set language]
  utm_content  = [uid_][ad name_][format_][ad language]
A parameter with no known part at all is left out. Source and medium are always known (platform / generator rows)."""
import re, sys, os, json, collections
from urllib.parse import urlsplit, urlunsplit, parse_qsl, quote
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from derive import PLATFORM_SOURCES, CLICK_IDS, KV_LETTER, MONTHS
from fallback import Fallback

CHANNEL_SRC = {'meta': 'meta', 'google': 'google', 'tiktok': 'tiktok', 'snapchat': 'snapchat', 'x': 'x', 'linkedin': 'linkedin'}
# objective from funnel + sub-objective words in a legacy campaign name: only combinations the generator's own
# Input_Campaign_Generator rows use for >= 90% of their rows (and at least 2 rows)
COMBO_OBJ = {('uf', 'reach'): 'awr', ('mf', 'reach'): 'awr', ('lf', 'reach'): 'awr', ('lf', 'leads'): 'conv', ('lf', 'sales'): 'conv',
             ('mf', 'traffic'): 'cond', ('lf', 'traffic'): 'cond'}
SUB = {'reach': 'reach', 'awareness': 'reach', 'leads': 'leads', 'lead': 'leads', 'sales': 'sales', 'conversion': 'sales', 'conversions': 'sales',
       'traffic': 'traffic'}

def generator_combos(gen_path):
    import openpyxl
    ws = openpyxl.load_workbook(gen_path, data_only=True)['Input_Campaign_Generator']
    H = [ws.cell(1, c).value for c in range(1, 26)]
    c = collections.defaultdict(collections.Counter)
    for r in range(3, ws.max_row + 1):
        d = dict(zip(H, [ws.cell(r, k).value for k in range(1, 26)]))
        if d.get('Campaign Funnel') and d.get('Campaign Objective') and d.get('Campaign Sub-Objective'):
            c[(d['Campaign Funnel'], d['Campaign Sub-Objective'])][d['Campaign Objective']] += 1
    out = {}
    for k, v in c.items():
        (o, n), = v.most_common(1); tot = sum(v.values())
        if tot >= 2 and n / tot >= 0.9: out[k] = o
    return out

class Omit:
    def __init__(self, deriver, scored, gen_path):
        self.D = deriver; self.N = deriver.N
        self.L = Fallback(deriver, scored, gen_path)
        self.combos = generator_combos(gen_path)

    # ---------- helpers ----------
    def objective(self, campaign_name):
        n = str(campaign_name or '')
        m = re.search(r'_(UF|MF|LF|UP)_(AWR|CONV|LEAD|TRF)_', n)
        if m: return {'AWR': 'awr', 'CONV': 'conv', 'LEAD': 'conv', 'TRF': 'cond'}[m.group(2)], f'objective from campaign-name token {m.group(2)}'
        f = re.search(r'(?:^|_)(UF|MF|LF|UP)(?:_|$)', n.upper())
        words = re.findall(r'[a-z]+', n.lower())
        subs = [SUB[w] for w in words if w in SUB]
        if f and subs:
            fun = {'UP': 'uf'}.get(f.group(1), f.group(1).lower())
            o = self.combos.get((fun, subs[0]))
            if o: return o, f'objective {o}: generator rows use {o} for funnel {fun} + sub-objective {subs[0]}'
        return None, None

    def compose(self, url, src, med, camp, term, content):
        parts = urlsplit(url.strip()); q = parse_qsl(parts.query, keep_blank_values=True)
        j = lambda xs: '_'.join(x for x in xs if x)
        exp = [('utm_source', src), ('utm_medium', med), ('utm_campaign', j(camp)), ('utm_term', j(term)), ('utm_content', j(content))]
        exp = [(k, v) for k, v in exp if v]
        keep = [(k, v) for k, v in q if not k.lower().startswith('utm_') and k.lower() not in CLICK_IDS]
        query = '&'.join(f'{k}={quote(v, safe="-_.~{}/:")}' for k, v in keep + exp)
        return urlunsplit((parts.scheme, parts.netloc, parts.path, query, parts.fragment))

    # ---------- legacy ads (old names) ----------
    def legacy(self, s, month):
        """month: 'mmmYY' (first serving) or None (never served -> left out)."""
        p = s['platform']
        g = self.L._pieces(s, month or 'x')
        omitted, basis = [], []
        src = g.get('src') or CHANNEL_SRC.get(p)
        med = g.get('med')
        if not med:
            meds = {r['medium'] for r in self.D.input_rows if r['source'] == src and r['medium']}
            med = meds.pop() if len(meds) == 1 else ('paid-search' if p == 'google' else 'paid-social')
            basis.append(f'utm_medium {med} (generator rows for {src})')
        obj = g.get('obj')
        if not obj:
            obj, why = self.objective(s.get('campaign_name'))
            if why: basis.append(why)
        cname = g.get('cname'); cuid = g.get('cuid') if cname else None
        tname = g.get('tname'); tuid = g.get('tuid') if tname else None
        aname = g.get('aname'); auid = g.get('auid') if aname else None
        tlang, afmt, alang = g.get('tlang'), g.get('afmt'), g.get('alang')
        lp = re.search(r'/(en|ar)(/|$)', urlsplit((s.get('url') or '').strip()).path.lower())
        if not tlang and lp: tlang = lp.group(1); basis.append(f'ad set language from landing page /{tlang}/')
        if not alang and lp: alang = lp.group(1); basis.append(f'ad language from landing page /{alang}/')
        if not afmt and p == 'google' and re.search(r'RESPONSIVE_SEARCH|channel=SEARCH|EXPANDED_TEXT', str(s.get('notes'))): afmt = 'search'
        for nm, v in (('campaign name', cname), ('month (ad never served)', month), ('objective', obj), ('ad set name', tname),
                      ('ad set language', tlang), ('ad name', aname), ('ad format', afmt), ('ad language', alang)):
            if not v: omitted.append(nm)
        url = self.compose(s['url'], src, med, [cuid, cname, month, obj], [tuid, tname, tlang], [auid, aname, afmt, alang])
        return url, omitted, basis

    # ---------- generator-format ads whose ad token is unknown (kva on social/search, Google search without a creative code) ----------
    def unknown_ad_token(self, s):
        p = s['platform']
        camp, ce = self.D.parse_campaign(s.get('campaign_name'))
        adset, ae = self.D.parse_adset(s.get('adset_name'))
        if not camp: return None, ['campaign not in generator format: ' + str(ce)], []
        cur = {k.lower(): v.lower() for k, v in parse_qsl(urlsplit(s['url'].strip()).query)}
        src = camp['src']
        if p in PLATFORM_SOURCES and src not in PLATFORM_SOURCES[p]: return None, [f'campaign source {src} not for {p}'], []
        med, mb = self.D.medium_for(src, camp['cname'], cur.get('utm_medium', ''))
        basis = ['ad token left out: ' + ('kva/kvb/kvd is programmatic-only' if p != 'google' or KV_LETTER.search(cur.get('utm_content', '')) else 'no search creative code in the generator')]
        mon, yy = camp['mon'], camp['year'][2:]
        fs = self.D.first_serving.get(str(s.get('ad_id')))
        if fs and s.get('_use_serving_month'): mon, yy = MONTHS[int(fs[5:7]) - 1], fs[2:4]
        term = [adset['gen_uid'], adset['aname'], adset['lang']] if adset else [None, None, None]
        omitted = []
        if not adset:
            # keep the live utm_term when it already is a valid generator pair + language (nothing invented)
            mt = re.fullmatch(r'([0-9a-f]{8})_([a-z0-9-]+)_(en-ar|ar|en|all)', cur.get('utm_term', ''))
            if mt and self.N.adset.get(mt.group(2)) == mt.group(1):
                term = [mt.group(1), mt.group(2), mt.group(3)]; basis.append('utm_term kept: already a valid generator ad set code + name + language')
            else:
                omitted.append('ad set (' + str(ae) + ')')
        # format + language of the ad: from the ad name tail, else the current content tail, else (Google) search + ad group language
        fm = '|'.join(sorted(self.N.formats, key=len, reverse=True))
        m = re.search(rf'_({fm})_(en-ar|ar|en|all)(?:_|$)', str(s.get('ad_name') or '').lower()) or re.search(rf'_({fm})_(en-ar|ar|en|all)$', cur.get('utm_content', ''))
        afmt, alang = (m.group(1), m.group(2)) if m else ('search' if p == 'google' else None, adset['lang'] if adset and p == 'google' else None)
        if not afmt: omitted.append('ad format')
        if not alang: omitted.append('ad language')
        url = self.compose(s['url'], src, med, [camp['gen_uid'], camp['cname'], f'{mon}{yy}', camp['obj']], term, [None, None, afmt, alang])
        return url, ['ad token (uid + ad name)'] + omitted, basis

def check_omit(N, url):
    """Validate a URL built under the omit rule: every part that IS present must be valid; missing parts are allowed."""
    errs = []
    parts = urlsplit(url.strip()); params = parse_qsl(parts.query, keep_blank_values=True)
    q = {}
    for k, v in params: q.setdefault(k, v)
    if re.search(r'x{4,}|\{[a-z_]+\}|(^|[_-])test([_-]|$)|adset_name|ad_name', url, re.I): errs.append('placeholder / macro left')
    if 'utm_id' in q: errs.append('utm_id present')
    extra = [k for k, _ in params if (k.startswith('utm_') and k not in ('utm_source', 'utm_medium', 'utm_campaign', 'utm_term', 'utm_content')) or k in CLICK_IDS]
    if extra: errs.append('extra params ' + ','.join(extra))
    for k in ('utm_source', 'utm_medium', 'utm_campaign', 'utm_term', 'utm_content'):
        v = q.get(k)
        if v is None: continue
        if v != v.lower(): errs.append(f'{k} upper-case')
        if not re.fullmatch(r'[a-z0-9_-]+', v): errs.append(f'{k} has characters outside a-z0-9_-')
    if q.get('utm_source', '') not in N.sources: errs.append('source not in Mappings')
    if q.get('utm_medium', '') not in N.mediums: errs.append('medium not in Mappings')
    def split(v, tails):
        toks = v.split('_'); got = {}
        for name, allowed in reversed(tails):
            if toks and allowed(toks[-1]): got[name] = toks.pop()
        if toks and re.fullmatch(r'[0-9a-f]{8}', toks[0]): got['uid'] = toks.pop(0)
        if toks: got['name'] = '_'.join(toks)
        return got
    mon = lambda t: re.fullmatch(r'(' + '|'.join(MONTHS) + r')\d\d', t) is not None
    for k, tails, lookup in (('utm_campaign', [('month', mon), ('objective', lambda t: t in N.objectives)], N.camp),
                             ('utm_term', [('lang', lambda t: t in N.langs)], N.adset),
                             ('utm_content', [('format', lambda t: t in N.formats), ('lang', lambda t: t in N.langs)], N.ad)):
        v = q.get(k)
        if not v: continue
        g = split(v, tails)
        if 'name' in g and '_' in g['name']: errs.append(f'{k}: unexpected "_" inside name "{g["name"]}"')
        if 'uid' in g and lookup.get(g.get('name', '')) != g['uid']: errs.append(f'{k}: uid {g["uid"]} does not belong to "{g.get("name")}" in Campaign_Code_Generator')
        if 'uid' not in g and g.get('name') in lookup: errs.append(f'{k}: generator has a code for "{g["name"]}" ({lookup[g["name"]]}) but it is missing')
        if k == 'utm_content' and g.get('name') and KV_LETTER.search(g['name']) and q.get('utm_medium') != 'programmatic': errs.append('kva token on social/search')
    return errs
